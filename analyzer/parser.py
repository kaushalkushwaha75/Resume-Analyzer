"""
Resume and document parser.
Extracts raw text from PDF and DOCX files.
"""

import os

import pdfplumber
from docx import Document


def parse_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file using pdfplumber.

    Args:
        file_path: Absolute path to the PDF file.

    Returns:
        Extracted text content as a single string.
    """
    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

                # Also attempt to extract text from tables
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        row_text = " | ".join(
                            cell.strip() if cell else "" for cell in row
                        )
                        if row_text.strip(" |"):
                            text_parts.append(row_text)
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {e}")

    if not text_parts:
        raise ValueError("No text could be extracted from the PDF. The file may be image-based or corrupted.")

    return "\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file using python-docx.

    Args:
        file_path: Absolute path to the DOCX file.

    Returns:
        Extracted text content as a single string.
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {e}")

    text_parts = []

    # Extract paragraph text
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                text_parts.append(row_text)

    if not text_parts:
        raise ValueError("No text could be extracted from the DOCX file.")

    return "\n".join(text_parts)


def parse_file(file_path: str) -> str:
    """
    Detect file type and route to the appropriate parser.

    Args:
        file_path: Absolute path to the resume file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If the file type is unsupported or parsing fails.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Please upload a PDF or DOCX file.")
